from io import BytesIO

import pytest
from django.core.exceptions import PermissionDenied, ValidationError
from docx import Document

from apps.evaluations import services
from apps.evaluations.models import Assessment, EvaluationTeamMember
from apps.reports.models import Report
from apps.reports.services import generate_report, report_availability
from apps.rules.systems import MRL, TRL
from tests.test_core_flow import fill, make_project


def _complete(assessment, owner, admin):
    for level in range(1, assessment.target_level + 1):
        fill(assessment, level, owner)
    services.submit_levels(assessment, owner, assessment.target_level)
    services.pass_levels(assessment, admin, assessment.target_level)


def _team(project):
    EvaluationTeamMember.objects.create(project=project, order=1, role="组长", name="杜工")


def _docx_text(report):
    document = Document(BytesIO(report.file.read()))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_trl_report_blocked_then_generated(enterprise, admin_user):
    project = make_project(enterprise, targets={TRL: 2})
    assessment = project.get_assessment(TRL)
    with pytest.raises(ValidationError):
        generate_report(project, Report.TYPE_TRL, admin_user)
    _complete(assessment, enterprise, admin_user)
    with pytest.raises(ValidationError):  # 缺评价组成员
        generate_report(project, Report.TYPE_TRL, admin_user)
    _team(project)
    with pytest.raises(PermissionDenied):
        generate_report(project, Report.TYPE_TRL, enterprise)

    report1 = generate_report(project, Report.TYPE_TRL, admin_user)
    assessment.refresh_from_db()
    assert assessment.status == Assessment.STATUS_REPORTED
    assert report1.version == 1 and report1.trl_level == 2 and report1.mrl_level is None
    assert report1.report_no.endswith("-TRL")
    text = _docx_text(report1)
    assert "技术就绪度评价报告" in text and "技术类型" in text and "杜工" in text
    assert "轨道" not in text and "Demo" not in text

    services.reopen_assessment(assessment, admin_user)
    report2 = generate_report(project, Report.TYPE_TRL, admin_user)
    assert report2.version == 2


def test_mrl_and_combined_reports(enterprise, admin_user, mrl_rules):
    project = make_project(enterprise, targets={TRL: 1, MRL: 2})
    _team(project)
    availability = report_availability(project)
    assert availability[Report.TYPE_COMBINED]["available"] and availability[Report.TYPE_COMBINED]["reasons"]

    _complete(project.get_assessment(MRL), enterprise, admin_user)
    mrl_report = generate_report(project, Report.TYPE_MRL, admin_user)
    assert mrl_report.mrl_level == 2 and mrl_report.report_no.endswith("-MRL")
    assert "制造成熟度评价报告" in _docx_text(mrl_report)

    with pytest.raises(ValidationError):  # TRL 尚未完成
        generate_report(project, Report.TYPE_COMBINED, admin_user)
    _complete(project.get_assessment(TRL), enterprise, admin_user)
    combined = generate_report(project, Report.TYPE_COMBINED, admin_user)
    assert combined.trl_level == 1 and combined.mrl_level == 2
    text = _docx_text(combined)
    assert "技术就绪度与制造成熟度综合评价报告" in text
    assert "4 技术就绪度评价" in text and "5 制造成熟度评价" in text and "6 总结" in text


def test_report_type_unavailable_without_assessment(enterprise, admin_user):
    project = make_project(enterprise, targets={TRL: 1})
    assert report_availability(project)[Report.TYPE_MRL]["available"] is False


def test_enterprise_can_download_own_report(client, enterprise, other_enterprise, admin_user):
    project = make_project(enterprise, targets={TRL: 1})
    _complete(project.get_assessment(TRL), enterprise, admin_user)
    _team(project)
    report = generate_report(project, Report.TYPE_TRL, admin_user)
    client.force_login(enterprise)
    assert client.get(f"/reports/{report.id}/download/").status_code == 200
    assert project.name in client.get("/reports/").content.decode()
    client.force_login(other_enterprise)
    assert client.get(f"/reports/{report.id}/download/").status_code == 404
    assert project.name not in client.get("/reports/").content.decode()


def test_report_preview_tabs_render(client, enterprise, admin_user, mrl_rules):
    project = make_project(enterprise, targets={TRL: 1, MRL: 1})
    client.force_login(admin_user)
    for report_type in ("trl", "mrl", "combined"):
        response = client.get(f"/reports/projects/{project.id}/?type={report_type}")
        assert response.status_code == 200
