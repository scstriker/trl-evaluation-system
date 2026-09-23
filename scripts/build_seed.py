# -*- coding: utf-8 -*-
"""生成评价细则种子 CSV（开发 / 维护期脚本，需要 openpyxl）。

用法（在 trl-evaluation-system/ 目录下执行）：
  python scripts/build_seed.py trl                  从上级目录 app_GUI.py 的常量生成 data/seed/trl_rules_seed.csv
  python scripts/build_seed.py mrl <细则.xlsx>       解析制造成熟度评价细则 xlsx，生成 data/seed/mrl_rules_seed.csv
  python scripts/build_seed.py mrl-template         生成制造成熟度评价细则填写模板 data/templates/制造成熟度评价细则-模板.xlsx

xlsx 版式与《评价细则.xlsx》一致：某一行含“序号”表头；其后每级一行“第 N 级 | 等级定义”，
下面逐条填写“序号 | 适用（硬件/软件/通用，可留空）| 具体化等级条件 | 评价支撑信息”。
生成 CSV 后执行 `python manage.py import_rules` 导入系统。
"""

import csv
import importlib.util
import re
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
APP_GUI = BASE_DIR.parent / "app_GUI.py"
SEED_DIR = BASE_DIR / "data" / "seed"
TEMPLATE_PATH = BASE_DIR / "data" / "templates" / "制造成熟度评价细则-模板.xlsx"
HEADER = ["row_type", "level", "track", "seq", "name", "evidence", "desc"]
LEVEL_RE = re.compile(r"^第\s*(\d+)\s*级$")
TRACKS = {"硬件": "hw", "软件": "sw", "通用": "gen", "": "gen"}


def load_app_gui():
    spec = importlib.util.spec_from_file_location("app_gui_source", APP_GUI)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def build_trl():
    module = load_app_gui()
    path = SEED_DIR / "trl_rules_seed.csv"
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(HEADER)
        for level in range(1, 10):
            writer.writerow(["level", level, "", "", module.LEVEL_NAMES[level], "", ""])
            for track, criteria in (("hw", module.HW_CRITERIA), ("sw", module.SW_CRITERIA)):
                for item in criteria.get(level, []):
                    writer.writerow(["item", level, track, item["seq"], item["name"], item["evidence"], ""])
    hw = sum(len(v) for v in module.HW_CRITERIA.values())
    sw = sum(len(v) for v in module.SW_CRITERIA.values())
    print(f"已生成 {path.name}（硬件 {hw} 条 + 软件 {sw} 条 = {hw + sw} 条）")


def _existing_level_desc(path):
    if not path.exists():
        return {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {int(r["level"]): r.get("desc") or "" for r in csv.DictReader(f) if r["row_type"] == "level"}


def _text(value):
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def parse_rules_xlsx(xlsx_path, max_level):
    from openpyxl import load_workbook

    sheet = load_workbook(xlsx_path, data_only=True).worksheets[0]
    rows = [[_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    header_index = next((i for i, row in enumerate(rows) if "序号" in row), None)
    if header_index is None:
        raise SystemExit("未找到含“序号”的表头行，请按模板版式整理。")
    col_seq = rows[header_index].index("序号")
    col_kind, col_text, col_evidence = col_seq + 1, col_seq + 2, col_seq + 3

    levels, items, errors = {}, [], []
    current, auto_seq = None, 0
    for number, row in enumerate(rows[header_index:], start=header_index + 1):
        row = row + [""] * (col_evidence + 1 - len(row))
        kind, text = row[col_kind], row[col_text]
        match = LEVEL_RE.match(kind)
        if number == header_index + 1 and not match:
            continue  # 纯表头行（《评价细则.xlsx》的表头同时是第 1 级行，需保留）
        if match:
            current, auto_seq = int(match.group(1)), 0
            if not 1 <= current <= max_level:
                errors.append(f"第 {number} 行：等级 {current} 超出 1~{max_level}")
            levels[current] = text
            continue
        if not text:
            continue
        if current is None:
            errors.append(f"第 {number} 行：条目出现在任何“第 N 级”之前")
            continue
        if kind not in TRACKS:
            errors.append(f"第 {number} 行：“适用”只能填硬件 / 软件 / 通用，或留空")
            continue
        auto_seq += 1
        seq = int(float(row[col_seq])) if re.fullmatch(r"\d+(\.0)?", row[col_seq]) else auto_seq
        items.append({"level": current, "track": TRACKS[kind], "seq": seq, "name": text, "evidence": row[col_evidence]})

    seen = set()
    for item in items:
        key = (item["level"], item["track"], item["seq"])
        if key in seen:
            errors.append(f"第 {item['level']} 级序号 {item['seq']} 重复")
        seen.add(key)
    missing = [level for level in range(1, max_level + 1) if level not in levels]
    if missing:
        errors.append(f"缺少等级行：第 {'、'.join(map(str, missing))} 级")
    if errors:
        raise SystemExit("细则文件有误：\n  " + "\n  ".join(errors))
    return levels, items


def build_mrl(xlsx_path):
    path = SEED_DIR / "mrl_rules_seed.csv"
    descs = _existing_level_desc(path)
    levels, items = parse_rules_xlsx(xlsx_path, max_level=10)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(HEADER)
        for level in sorted(levels):
            writer.writerow(["level", level, "", "", levels[level], "", descs.get(level, "")])
            for item in (i for i in items if i["level"] == level):
                writer.writerow(["item", level, item["track"], item["seq"], item["name"], item["evidence"], ""])
    print(f"已生成 {path.name}：{len(levels)} 个等级、{len(items)} 条细则。下一步执行 python manage.py import_rules")


def build_mrl_template():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    with open(SEED_DIR / "mrl_rules_seed.csv", encoding="utf-8-sig", newline="") as f:
        names = {int(r["level"]): r["name"] for r in csv.DictReader(f) if r["row_type"] == "level"}
    wb = Workbook()
    ws = wb.active
    ws.title = "制造成熟度评价细则"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    level_fill = PatternFill("solid", fgColor="DDE7F0")
    ws.append(["", "制造成熟度等级", "序号", "适用", "具体化等级条件", "评价支撑信息"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for level in range(1, 11):
        ws.append(["", "", "", f"第{level}级", names.get(level, ""), ""])
        for cell in ws[ws.max_row]:
            cell.fill = level_fill
            cell.font = Font(bold=True)
        for seq in range(1, 4):
            ws.append(["", "", seq, "通用", "", ""])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=2, max_col=6):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    for column, width in zip("ABCDEF", (2, 14, 6, 8, 60, 28)):
        ws.column_dimensions[column].width = width

    guide = wb.create_sheet("填写说明")
    for line in [
        "1. 每个等级保留一行“第 N 级”，在 E 列填写等级定义（已按现有定义预填，可修改）。",
        "2. 在该等级行下方逐条填写：C 列序号、D 列适用（通用 / 硬件 / 软件，一般填“通用”）、E 列具体化等级条件、F 列评价支撑信息。",
        "3. 条目数量不限，可插入或删除行；E 列为空的行会被忽略。",
        "4. 同一等级内序号不可重复；MRL 1~10 级都需要有等级行。",
        "5. 整理完成后交评价机构管理员，执行 python scripts/build_seed.py mrl <本文件> 与 python manage.py import_rules 导入。",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TEMPLATE_PATH)
    print(f"已生成模板 {TEMPLATE_PATH.relative_to(BASE_DIR)}")


def main(argv):
    command = argv[1] if len(argv) > 1 else "trl"
    if command == "trl":
        build_trl()
    elif command == "mrl" and len(argv) > 2:
        build_mrl(Path(argv[2]))
    elif command == "mrl-template":
        build_mrl_template()
    else:
        print(__doc__)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
