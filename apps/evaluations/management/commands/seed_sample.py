# -*- coding: utf-8 -*-
"""初始化数据：管理员账号 + 样例企业账号与各状态的评价项目（口径参照《4400 技术就绪度评价报告》）。

用法：python manage.py seed_sample            # 管理员 + 样例企业与项目（可重复执行，会重建样例项目）
      python manage.py seed_sample --admin-only  # 仅创建 / 重置管理员账号
"""

from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import BaseCommand
from django.db import transaction
from django.utils import timezone
from docx import Document

from apps.accounts.forms import credit_code_check_char
from apps.accounts.models import EnterpriseProfile
from apps.evaluations import services
from apps.evaluations.models import CheckItem, EvaluationProject, EvaluationTeamMember
from apps.evidence.services import upload_evidence_file
from apps.reports.models import Report
from apps.reports.services import generate_report
from apps.rules.models import RuleItem
from apps.rules.systems import TRL

INITIAL_PASSWORD = "Trl@2026"

ADMIN_USERNAME = "admin"

# (用户名, 企业全称, 信用代码前 17 位, 联系人, 手机, 状态)
ENTERPRISES = [
    ("gangyan", "钢研纳克检测技术股份有限公司", "91110000MA0000001", "张工", "13800000001", EnterpriseProfile.STATUS_ACTIVE),
    ("hexin", "核芯光电科技（山东）有限公司", "91370100MA0000002", "李工", "13800000002", EnterpriseProfile.STATUS_ACTIVE),
    ("zhongnan", "中南大学", "12430000MB0000003", "王老师", "13800000003", EnterpriseProfile.STATUS_ACTIVE),
    ("jingce", "江苏精测智能装备有限公司", "91320500MA0000004", "赵经理", "13800000004", EnterpriseProfile.STATUS_PENDING),
]

TEAM = [
    ("组长", "杜工", "正高级工程师", "测量控制设备及系统评价", "机械工业仪器仪表综合技术经济研究所"),
    ("成员", "李工", "高级工程师", "智能制造与装备技术", "机械工业仪器仪表综合技术经济研究所"),
    ("成员", "王工", "高级工程师", "工业软件测评", "机械工业仪器仪表综合技术经济研究所"),
]

OVERVIEW_TEXT = (
    "本项目来源于“十五五”智能制造系统和机器人国家科技重大专项，由钢研纳克检测技术股份有限公司牵头，"
    "针对稀土萃取分离过程成分离线检测时间长、数据获取滞后、工艺参数在线感知不及时等问题，"
    "突破网格化采样、多流路全自动样品传送、严苛环境可靠性设计等关键技术，"
    "研制高能量高稳定激发源、高效降噪光路系统和多参量同步采传模块等核心部件，"
    "开发具备多流路状态监控、流路识别与方法切换协同控制、谱图解算和数据交互等功能的一体化软件，"
    "研制稀土萃取分离过程元素成分在线检测仪表，在稀土、钨钼等湿法冶金场景中完成仪表性能验证。\n"
    "主要技术指标：激发源电压≥65kV、功率≥65W；SDD 探测器分辨率＜136eV@5.9keV；"
    "稀土元素检出限≤0.01%，检测重复性 RSD≤1%，在线检测时间≤2min；整机技术就绪度（TRL）≥8 级。"
)

SUMMARY_TEXT = (
    "本项目聚焦稀土萃取分离过程元素成分在线检测仪表的国产化研制，围绕取样系统、高稳定性光源、"
    "高性能光学系统、加密数据互传系统等关键核心部件，统筹推进方案设计、部件研制、软件开发、整机集成与性能验证。"
    "在设计上已有详细的设计方案及实施方案，并形成了相关专利；试制上已形成初样机，满足设计要求且持续优化中；"
    "部分关键部件已通过第三方检测机构验证。"
)

# 各级“满足情况说明 / 评价支撑信息”样例文本（按 TRL 级别 + 类别 + 序号）
STATEMENTS = {
    (1, "hw", 1): ("明确了 X 射线荧光分析系统、数据处理系统和自动控制系统的完整技术链路，核心原理已通过发明专利公开，为硬件研发奠定理论基础。", "发明专利《一种稀土冶炼分离过程质量配分量在线监测仪》"),
    (1, "hw", 2): ("项目任务书系统界定了仪表的适用环境与检测范围，论文完成原理假设的边界论证，明确技术的适用场景与约束条件。", "《项目任务书》；论文《基于能量色散-X射线荧光光谱方法对轻稀土料液配分含量的在线测定》"),
    (1, "sw", 3): ("项目任务书已完成软件系统全链路技术瓶颈梳理，识别出多流路控制、机器人协同、仪器状态监控、谱数据处理与传输等关键问题。", "《项目任务书》"),
    (1, "sw", 4): ("基于能量色散 X 射线荧光光谱技术建立了在线测定稀土配分含量的方法，核心算法原理已通过论文公开。", "论文《X射线荧光光谱法在线测定稀土冶炼分离过程中钬铒铥镱》"),
    (1, "sw", 5): ("明确了核心算法的运行条件与适用范围，通过整体技术路线可行性论证确认软件系统研发的技术可行性。", "《软件可行性报告》《机器人辅助取样系统开发文档》"),
    (2, "hw", 1): ("任务书明确仪表核心性能指标，完成硬件系统初步实施方案编制，梳理核心硬件模块的技术要素与结构特性。", "《项目任务书》《项目实施工作方案》"),
    (2, "hw", 2): ("初步明确硬件系统可实现的核心功能（多点位取样、在线检测、数据传输），确定了定量指标。", "《项目任务书》"),
    (2, "hw", 3): ("明确仪表在高温、高湿、高酸、强有机腐蚀等萃取分离现场环境下工作的预期应用环境要求。", "《项目任务书》《应用环境分析文档》"),
    (2, "sw", 4): ("完成一体化软件系统需求分析，形成多流路状态监控、谱图解算、数据交互等功能需求清单。", "《软件需求分析文档》"),
    (2, "sw", 5): ("确定基于工业现场数据与机理知识融合的一体化智能分析技术路线。", "《软件设计文档》"),
    (2, "sw", 6): ("完成光谱预处理、重叠峰解析、元素检测等模块的技术准备与预研。", "《软件技术准备报告》"),
    (2, "sw", 7): ("形成一体化软件概要设计，划分自检、采样、检测、数据交互、可视化五大模块。", "《概要设计说明书》"),
    (3, "hw", 1): ("形成两级管道中继取样、机器人辅助取样、X 射线源、XRF 光学、无线加密互传五大子系统的完善实施方案。", "《项目实施工作方案》"),
    (3, "hw", 2): ("通过仿真与台架试验验证了多流路取样、高压稳定激发等关键功能的可行性。", "《关键功能仿真验证报告》"),
    (3, "hw", 3): ("理论分析了光管、探测器、无线系统与整机的集成方案可行性。", "《系统集成可行性分析文档》"),
    (3, "hw", 4): ("形成覆盖五个课题的项目开发计划，明确里程碑与交付物。", "《项目开发计划》"),
    (3, "hw", 5): ("评估了 X 射线管真空封接、SDD 芯片封装等制造条件与现有制造能力。", "《制造能力评估报告》"),
    (3, "sw", 6): ("确定软件需求边界：覆盖 17 种稀土元素与 6 种以上混合料液检测。", "《需求边界确认报告》"),
    (3, "sw", 7): ("完成解谱算法等关键技术验证，软件可对稀土元素谱峰进行识别。", "《软件关键技术验证报告》"),
    (3, "sw", 8): ("完成一体化软件详细设计。", "《软件详细设计文档》"),
    (4, "hw", 1): ("完成两级管道中继取样系统、X 射线源等关键功能试样开发，形成专利与研究报告并被采纳。", "《关键部件开发文档》；发明专利受理通知书"),
    (4, "hw", 2): ("在实验室环境下完成光管 50kV/50W 工况测试、探测器分辨率 176eV@5.9keV 测试。", "《大功率光管测试报告》《高性能XRF光学系统检测报告》"),
    (4, "hw", 3): ("试制了两级管道中继取样系统样件、机器人初样机、X 射线管样机、SDD 探头初样机。", "实物样机照片；《试制记录》"),
    (4, "hw", 4): ("集成光管、探测器、无线加密系统，形成 15 路测试通道的光谱仪初样机。", "《整机集成报告》"),
    (4, "hw", 5): ("评估了真空焊接、无尘装配、SDD 芯片筛选等关键制造工艺。", "《大功率X射线管制造工艺》《高性能XRF光学系统制备工艺方案》"),
    (4, "hw", 6): ("各关键功能试样设计过程文档清晰，可追溯。", "《设计过程文档清单》"),
    (4, "sw", 7): ("完成软件研发实施方案及进度计划。", "《软件实施方案及进度计划文档》"),
    (4, "sw", 8): ("完成软件主框架研发，原型系统在初样机上运行。", "《软件技术路线文档》"),
    (4, "sw", 9): ("基于原型系统开展全流程功能验证、性能测试与兼容性分析。", "《软件关键技术验证报告》"),
}

DEFAULT_STATEMENT = "经评价组核查相关佐证材料，该具体化等级条件已满足。"


def _evidence_docx(title):
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph("本文件为系统初始化时生成的样例佐证材料，正式评价时请上传真实文件。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class Command(BaseCommand):
    help = "初始化管理员账号与样例企业、样例评价项目（可重复执行，会重建样例企业的项目）"

    def add_arguments(self, parser):
        parser.add_argument("--admin-only", action="store_true", help="仅创建 / 重置管理员账号，不生成样例企业与项目")

    @transaction.atomic
    def handle(self, *args, **options):
        if not RuleItem.objects.filter(system=TRL).exists():
            call_command("import_rules", verbosity=0)

        admin, _ = User.objects.get_or_create(username=ADMIN_USERNAME)
        admin.set_password(INITIAL_PASSWORD)
        admin.first_name = "评价中心管理员"
        admin.is_staff = admin.is_superuser = admin.is_active = True
        admin.save()
        self.stdout.write(self.style.SUCCESS(f"管理员账号：{ADMIN_USERNAME}（初始密码 {INITIAL_PASSWORD}，请登录后修改）"))
        if options["admin_only"]:
            return

        users = {}
        for username, company, code17, contact, phone, status in ENTERPRISES:
            user, _ = User.objects.get_or_create(username=username)
            user.set_password(INITIAL_PASSWORD)
            user.first_name = contact
            user.is_active = status == EnterpriseProfile.STATUS_ACTIVE
            user.save()
            EnterpriseProfile.objects.update_or_create(
                user=user,
                defaults={
                    "company_name": company,
                    "credit_code": code17 + credit_code_check_char(code17),
                    "contact_name": contact,
                    "phone": phone,
                    "status": status,
                    "reviewed_by": admin if status == EnterpriseProfile.STATUS_ACTIVE else None,
                    "reviewed_at": timezone.now() if status == EnterpriseProfile.STATUS_ACTIVE else None,
                },
            )
            users[username] = user
        EvaluationProject.objects.filter(owner__in=users.values()).delete()
        self.stdout.write(self.style.SUCCESS(f"样例企业账号：{', '.join(users)}（初始密码 {INITIAL_PASSWORD}；jingce 为待审核注册）"))

        gangyan, hexin, zhongnan = users["gangyan"], users["hexin"], users["zhongnan"]

        # ① 填写中：已通过 1~3 级，第 4 级部分填写
        p1 = self._create(
            gangyan,
            name="稀土萃取分离过程元素成分在线检测仪表",
            domain="稀土湿法冶金 · 在线检测仪表",
            tech_type=EvaluationProject.TECH_EQUIPMENT,
            target=8,
        )
        self._fill(p1, gangyan, levels=[1, 2, 3])
        self._submit_and_pass(p1, gangyan, admin, up_to=3)
        self._fill_partial(p1, gangyan, level=4, count=5)

        # ② 待审核：第 1~4 级一次提交
        p2 = self._create(
            gangyan,
            name="稀土萃取过程多流路在线取样系统",
            domain="稀土湿法冶金 · 取样与预处理",
            tech_type=EvaluationProject.TECH_HARDWARE,
            target=4,
        )
        self._fill(p2, gangyan, levels=[1, 2, 3, 4])
        services.submit_levels(p2.get_assessment(TRL), gangyan, 4)

        # ③ 待出报告：目标 4 级全部通过，评价组已填写
        p3 = self._create(
            hexin,
            name="高性能 XRF 光学系统",
            domain="X 射线探测器 · 核心器件",
            tech_type=EvaluationProject.TECH_HARDWARE,
            target=4,
        )
        self._fill(p3, hexin, levels=[1, 2, 3, 4])
        self._submit_and_pass(p3, hexin, admin, up_to=4)
        self._team(p3)

        # ④ 已退回：第 1 级通过，第 2~3 级退回修改（含逐条审核意见）
        p4 = self._create(
            hexin,
            name="大功率 X 射线管",
            domain="X 射线源 · 核心部件",
            tech_type=EvaluationProject.TECH_HARDWARE,
            target=3,
        )
        self._fill(p4, hexin, levels=[1, 2, 3])
        assessment = p4.get_assessment(TRL)
        services.submit_levels(assessment, hexin, 3)
        services.pass_levels(assessment, admin, 1)
        item = assessment.items.filter(level=2).order_by("track", "seq").first()
        services.set_review_comment(item, admin, "请补充方案评审会议纪要或专家评审意见，作为技术方案已确定的佐证。")
        services.return_levels(assessment, admin, [2, 3], "第 2 级技术方案的佐证不充分，请按条目意见补充材料后，与第 3 级一并重新提交。")

        # ⑤ 已出报告：软件 · 目标 3 级 · 已出具技术就绪度评价报告 V1
        p5 = self._create(
            zhongnan,
            name="稀土萃取在线检测一体化软件",
            domain="工业软件 · 光谱数据解析",
            tech_type=EvaluationProject.TECH_SOFTWARE,
            target=3,
        )
        self._fill(p5, zhongnan, levels=[1, 2, 3])
        self._submit_and_pass(p5, zhongnan, admin, up_to=3)
        self._team(p5)
        generate_report(p5, Report.TYPE_TRL, admin)

        self.stdout.write(self.style.SUCCESS("样例项目已生成：填写中 / 待审核 / 待出报告 / 已退回 / 已出报告 各一个"))

    # ------------------------------------------------------------------
    def _create(self, owner, *, name, domain, tech_type, target):
        project = EvaluationProject(
            name=name,
            applicant=owner.enterprise.company_name,
            domain=domain,
            tech_type=tech_type,
            overview_text=OVERVIEW_TEXT,
        )
        services.create_project(project, owner, {TRL: target})
        assessment = project.get_assessment(TRL)
        assessment.summary_text = SUMMARY_TEXT
        assessment.save(update_fields=["summary_text"])
        return project

    def _team(self, project):
        for order, (role, name, title, specialty, org) in enumerate(TEAM, start=1):
            EvaluationTeamMember.objects.create(
                project=project, order=order, role=role, name=name, title=title, specialty=specialty, org=org
            )

    def _fill_item(self, item, owner):
        statement, support = STATEMENTS.get((item.level, item.track, item.seq), (DEFAULT_STATEMENT, ""))
        title = support.split("；")[0].strip("《》") if support else f"{item.code} 佐证材料"
        upload_evidence_file(
            item,
            owner,
            SimpleUploadedFile(
                f"{title[:40]}.docx",
                _evidence_docx(title),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
        services.evaluate_check_item(
            item, owner, {"result": CheckItem.RESULT_SATISFIED, "statement": statement, "support_info": support}
        )

    def _fill(self, project, owner, levels):
        assessment = project.get_assessment(TRL)
        for item in assessment.items.filter(level__in=levels).order_by("level", "track", "seq"):
            self._fill_item(item, owner)

    def _fill_partial(self, project, owner, level, count):
        assessment = project.get_assessment(TRL)
        for item in assessment.items.filter(level=level).order_by("track", "seq")[:count]:
            self._fill_item(item, owner)

    def _submit_and_pass(self, project, owner, admin, up_to):
        assessment = project.get_assessment(TRL)
        services.submit_levels(assessment, owner, up_to)
        services.pass_levels(assessment, admin, up_to)
