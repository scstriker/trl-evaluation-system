"""评价报告：技术就绪度 / 制造成熟度 / 两者综合三种类型。

报告正文先装配为统一的“块”结构（标题、段落、表题、表格），在线预览与 Word 导出共用同一份数据，
保证两者内容一致。单体系报告沿用《4400 技术就绪度评价报告》的章节结构。
"""

import hashlib
from io import BytesIO

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Max
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from apps.accounts.permissions import can_generate_report, display_name
from apps.audit.models import AuditLog
from apps.evaluations.models import Assessment, CheckItem
from apps.evaluations.services import get_achieved_level, get_reviews
from apps.evidence.models import EvidenceFile
from apps.reports.models import Report
from apps.rules.models import LevelDefinition
from apps.rules.systems import MRL, TRL, get_system

RESULT_LABELS = {
    CheckItem.RESULT_SATISFIED: "是",
    CheckItem.RESULT_NOT_SATISFIED: "否",
    CheckItem.RESULT_NOT_APPLICABLE: "不适用",
    CheckItem.RESULT_UNEVALUATED: "—",
}

TYPE_SYSTEMS = {
    Report.TYPE_TRL: [TRL],
    Report.TYPE_MRL: [MRL],
    Report.TYPE_COMBINED: [TRL, MRL],
}

TYPE_SUFFIX = {Report.TYPE_TRL: "TRL", Report.TYPE_MRL: "MRL", Report.TYPE_COMBINED: "ZH"}

TYPE_TITLES = dict(Report.TYPE_CHOICES)


# ---------------------------------------------------------------------------
# 数据装配
# ---------------------------------------------------------------------------


def _assessment_part(assessment):
    sys = assessment.sys
    reviews = get_reviews(assessment)
    achieved = get_achieved_level(assessment, reviews)
    level_defs = {d.level: d.name for d in LevelDefinition.objects.filter(system=assessment.system)}
    items = list(
        assessment.items.filter(level__lte=achieved)
        .prefetch_related("evidence_files")
        .order_by("level", "track", "seq")
    )
    sections = []
    for level in range(1, achieved + 1):
        sections.append(
            {
                "level": level,
                "definition": level_defs.get(level, ""),
                "items": [item for item in items if item.level == level],
                "review": reviews.get(level),
            }
        )
    counts = {"completed": len(items), "satisfied": 0, "not_satisfied": 0, "not_applicable": 0}
    for item in items:
        if item.result in counts:
            counts[item.result] += 1
    return {
        "assessment": assessment,
        "sys": sys,
        "achieved": achieved,
        "target": assessment.target_level,
        "level_defs": level_defs,
        "sections": sections,
        "counts": counts,
        "has_general": any(item.track == CheckItem.TRACK_GENERAL for item in items),
    }


def report_availability(project):
    """每种报告能否生成及原因（报告页页签使用）。"""
    assessments = {a.system: a for a in project.assessments.all()}
    result = {}
    for report_type, systems in TYPE_SYSTEMS.items():
        reasons, missing = [], False
        for system in systems:
            assessment = assessments.get(system)
            if assessment is None:
                missing = True
                reasons.append(f"该项目未开展{get_system(system).name}评价")
                continue
            achieved = get_achieved_level(assessment)
            if achieved < assessment.target_level:
                sys = assessment.sys
                reasons.append(
                    f"{sys.name}目标等级 {sys.abbr} {assessment.target_level} 尚未全部审核通过（当前达成 {sys.abbr} {achieved}）"
                )
        result[report_type] = {"available": not missing, "reasons": reasons}
    return result


def build_report(project, report_type):
    systems = TYPE_SYSTEMS[report_type]
    assessments = {a.system: a for a in project.assessments.all()}
    parts = [_assessment_part(assessments[system]) for system in systems if system in assessments]
    if not parts:
        raise ValidationError("该项目没有可用于此类报告的评价。")
    team = list(project.team_members.all())
    evidence_files = list(
        EvidenceFile.objects.filter(check_item__assessment__in=[part["assessment"] for part in parts], is_deleted=False)
        .select_related("check_item__assessment")
        .order_by("check_item__assessment__system", "check_item__order_index", "uploaded_at")
    )
    title = TYPE_TITLES[report_type]
    blocks = _body_blocks(project, report_type, parts, team, evidence_files)
    return {
        "project": project,
        "report_type": report_type,
        "title": title,
        "parts": parts,
        "team": team,
        "evidence_files": evidence_files,
        "blocks": blocks,
    }


def _basis(parts):
    seen, rows = set(), []
    for part in parts:
        for code, name in part["sys"].basis:
            if name not in seen:
                seen.add(name)
                rows.append([code, name])
    return rows


def _levels_text(parts):
    return "、".join(f"{p['sys'].name} {p['achieved']} 级（{p['sys'].abbr} {p['achieved']}）" for p in parts)


def _criteria_blocks(part, chapter, table_no):
    """评价准则：等级定义表 + 已达成各级的条件表。返回 (blocks, 下一个表号)。"""
    sys = part["sys"]
    blocks = [
        {"kind": "caption", "text": f"表 {chapter}-{table_no} {sys.name}等级定义"},
        {
            "kind": "table",
            "headers": ["等级", f"{sys.name}等级定义"],
            "rows": [[level, part["level_defs"].get(level, "")] for level in sys.levels],
            "widths": [1.5, 14.5],
        },
    ]
    table_no += 1
    for section in part["sections"]:
        blocks.append({"kind": "caption", "text": f"表 {chapter}-{table_no} {sys.name} {section['level']} 级条件"})
        if part["has_general"]:
            headers = ["序号", f"第{section['level']}级：{section['definition']}", "评价支撑信息"]
            rows = [[item.seq, item.name, item.evidence] for item in section["items"]]
            widths = [1.2, 10.8, 4.0]
        else:
            headers = ["序号", "适用", f"第{section['level']}级：{section['definition']}", "评价支撑信息"]
            rows = [[item.seq, item.get_track_display(), item.name, item.evidence] for item in section["items"]]
            widths = [1.2, 1.4, 9.4, 4.0]
        blocks.append({"kind": "table", "headers": headers, "rows": rows, "widths": widths})
        table_no += 1
    return blocks, table_no


def _support_text(item):
    if item.support_info:
        return item.support_info
    return "、".join(f"《{e.original_filename}》" for e in item.evidence_files.all() if not e.is_deleted)


def _analysis_blocks(project, part, chapter):
    sys = part["sys"]
    assessment = part["assessment"]
    achieved = part["achieved"]
    blocks = [
        {"kind": "h2", "text": f"{chapter}.1 综述"},
    ]
    if assessment.summary_text:
        blocks.append({"kind": "p", "text": assessment.summary_text})
    blocks.append(
        {
            "kind": "p",
            "text": f"经评价组综合评估，“{project.name}”达到了{sys.standard}规定的{sys.name} {achieved} 级（{sys.abbr} {achieved}）的条件。",
        }
    )
    blocks.append({"kind": "h2", "text": f"{chapter}.2 确定等级的分析说明"})
    if not part["sections"]:
        blocks.append({"kind": "p", "text": "尚无审核通过的级别，逐级分析说明将在各级审核通过后生成。"})
        return blocks
    blocks.append({"kind": "p", "text": f"{sys.name}等级分析及确定见表 {chapter}-1 至表 {chapter}-{achieved}。"})
    for section in part["sections"]:
        review = section["review"]
        passed = f" · {timezone.localtime(review.reviewed_at):%Y-%m-%d} 审核通过" if review and review.reviewed_at else ""
        blocks.append(
            {"kind": "caption", "text": f"表 {chapter}-{section['level']} {sys.name}等级结论的分析说明（{section['level']} 级）{passed}"}
        )
        if part["has_general"]:
            headers = ["序号", "是否满足", "具体化等级条件", "满足情况说明", "评价支撑信息"]
            rows = [
                [item.seq, RESULT_LABELS[item.result], item.name, item.statement, _support_text(item)]
                for item in section["items"]
            ]
            widths = [1.0, 1.5, 5.0, 5.5, 3.0]
        else:
            headers = ["序号", "适用", "是否满足", "具体化等级条件", "满足情况说明", "评价支撑信息"]
            rows = [
                [item.seq, item.get_track_display(), RESULT_LABELS[item.result], item.name, item.statement, _support_text(item)]
                for item in section["items"]
            ]
            widths = [1.0, 1.2, 1.5, 4.3, 5.0, 3.0]
        blocks.append({"kind": "table", "headers": headers, "rows": rows, "widths": widths})
    return blocks


def _body_blocks(project, report_type, parts, team, evidence_files):
    combined = report_type == Report.TYPE_COMBINED
    names = "和".join(part["sys"].name for part in parts)
    blocks = []

    # 1 概况
    info_rows = [
        ["技术项目名称", project.name],
        ["委托（受评）单位", project.applicant],
        ["评价出具单位", project.issuer],
        ["应用领域", project.domain or "—"],
        ["技术类型", project.get_tech_type_display()],
    ]
    for part in parts:
        sys = part["sys"]
        info_rows.append([f"目标{sys.name}等级", f"{sys.abbr} {part['target']}"])
        info_rows.append([f"评价达成{sys.name}等级", f"{sys.abbr} {part['achieved']}"])
    blocks += [
        {"kind": "h1", "text": "1 概况"},
        {"kind": "h2", "text": "1.1 项目基本信息"},
        {"kind": "table", "headers": ["项目", "内容"], "rows": info_rows, "widths": [4.5, 11.5], "info": True},
        {"kind": "h2", "text": "1.2 项目概况"},
        {"kind": "p", "text": project.overview_text or "（项目概况由委托单位提供，参见项目任务书及研制总结材料。）"},
    ]

    # 2 评价实施情况
    targets = "；".join(f"{p['sys'].name} {p['target']} 级（{p['sys'].abbr} {p['target']}）" for p in parts)
    scope = "；".join(f"{p['sys'].abbr} 1 级至 {p['sys'].abbr} {p['achieved']} 级" for p in parts)
    standards = "、".join(p["sys"].standard for p in parts)
    blocks += [
        {"kind": "h1", "text": "2 评价实施情况"},
        {"kind": "h2", "text": "2.1 评价目的和范围"},
        {
            "kind": "p",
            "text": f"对{project.applicant}研制的“{project.name}”开展{names}评价，验证其是否满足任务书中规定的{targets}的要求。"
            f"评价范围覆盖{scope}的全部评价细则条目。",
        },
        {"kind": "h2", "text": "2.2 评价依据"},
        {"kind": "table", "headers": ["标准编号", "标准名称"], "rows": _basis(parts), "widths": [4.0, 12.0]},
        {"kind": "h2", "text": "2.3 评价人员"},
        {
            "kind": "table",
            "headers": ["角色", "姓名", "职称", "专业及特长", "工作单位"],
            "rows": [[m.role, m.name, m.title, m.specialty, m.org] for m in team] or [["—", "（待填写）", "", "", ""]],
            "widths": [1.6, 2.0, 2.6, 4.4, 5.4],
        },
        {"kind": "h2", "text": "2.4 评价过程"},
        {
            "kind": "p",
            "text": "评价工作按照“受理申请—企业自评与佐证提交—逐级审核—报告编制—审核批准”的流程实施。"
            f"委托单位依据{standards}对各级评价细则逐条自评并提交佐证材料，评价组逐级核查佐证材料并作出审核结论，"
            "每一级经审核通过后方纳入达成等级。全过程填写、提交、审核记录与佐证材料均在系统中留痕可溯。",
        },
    ]

    # 3 评价准则
    blocks.append({"kind": "h1", "text": "3 评价准则"})
    table_no = 1
    for index, part in enumerate(parts, start=1):
        if combined:
            blocks.append({"kind": "h2", "text": f"3.{index} {part['sys'].name}评价准则"})
        part_blocks, table_no = _criteria_blocks(part, 3, table_no)
        blocks += part_blocks

    # 4（及 5）等级评价
    chapter = 4
    for part in parts:
        blocks.append({"kind": "h1", "text": f"{chapter} {part['sys'].name}评价"})
        blocks += _analysis_blocks(project, part, chapter)
        chapter += 1

    # 总结
    blocks.append({"kind": "h1", "text": f"{chapter} 总结"})
    for part in parts:
        sys, counts = part["sys"], part["counts"]
        passed_dates = "；".join(
            f"{sys.abbr} {s['level']} 级于 {timezone.localtime(s['review'].reviewed_at):%Y-%m-%d} 审核通过"
            for s in part["sections"]
            if s["review"] and s["review"].reviewed_at
        )
        text = (
            f"{sys.name}评价累计核验细则条目 {counts['completed']} 条，其中自评满足 {counts['satisfied']} 条、"
            f"不满足 {counts['not_satisfied']} 条、不适用 {counts['not_applicable']} 条。"
        )
        if passed_dates:
            text += f"{passed_dates}。"
        blocks.append({"kind": "p", "text": text})
    blocks.append({"kind": "p", "text": f"综上所述，“{project.name}”的{_levels_text(parts)}。", "strong": True})

    # 附录A
    code_col = "关联细则"
    blocks += [
        {"kind": "h1", "text": "附录A 评价支撑信息清单"},
        {
            "kind": "table",
            "headers": ["编号", "评价支撑信息（佐证材料）", code_col, "SHA-256（前 16 位）"],
            "rows": [
                [index, e.original_filename, e.check_item.code, e.sha256[:16]]
                for index, e in enumerate(evidence_files, start=1)
            ]
            or [["—", "（无佐证材料）", "—", "—"]],
            "widths": [1.2, 8.3, 2.8, 3.7],
        },
    ]
    return blocks


# ---------------------------------------------------------------------------
# 生成条件
# ---------------------------------------------------------------------------


def generate_blockers(project, report_type, actor=None):
    availability = report_availability(project)[report_type]
    blockers = list(availability["reasons"])
    if not project.team_members.exists():
        blockers.append("评价组成员尚未填写（在“编辑报告信息”中填写）")
    if actor is not None and not can_generate_report(actor):
        blockers.append("报告由评价机构管理员生成")
    return blockers


def validate_can_generate(project, report_type, actor):
    if not can_generate_report(actor):
        raise PermissionDenied("报告由评价机构管理员生成。")
    if report_type not in TYPE_SYSTEMS:
        raise ValidationError("报告类型无效。")
    blockers = generate_blockers(project, report_type)
    if blockers:
        raise ValidationError(f"暂不能生成{TYPE_TITLES[report_type]}：{'；'.join(blockers)}。")


# ---------------------------------------------------------------------------
# Word 渲染
# ---------------------------------------------------------------------------


def _cn_year_month(dt):
    digits = "〇一二三四五六七八九"
    year = "".join(digits[int(ch)] for ch in str(dt.year))
    months = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十", "十一", "十二"]
    return f"{year}年{months[dt.month - 1]}月"


def _set_base_style(document):
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for level, size in ((1, 16), (2, 13)):
        heading = document.styles[f"Heading {level}"]
        heading.font.name = "Times New Roman"
        heading.font.size = Pt(size)
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def _para(document, text, *, align=None, bold=False, size=None):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def _add_table(document, headers, rows, *, header_bold=True, widths=None):
    """widths：各列宽度（cm），总宽约 16cm（A4 默认页边距下的正文宽度）。"""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        if header_bold:
            for run in cell.paragraphs[0].runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = "" if value is None else str(value)
    if widths:
        table.autofit = False
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    return table


def _docx_bytes(report, report_no, version, generated_by):
    project = report["project"]
    now = timezone.localtime()
    document = Document()
    _set_base_style(document)

    # ---------- 封面 ----------
    _para(document, f"报告编号：{report_no}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    for _ in range(4):
        document.add_paragraph()
    _para(document, report["title"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=24)
    document.add_paragraph()
    _para(document, project.name, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16)
    for _ in range(6):
        document.add_paragraph()
    _para(document, project.issuer, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    _para(document, _cn_year_month(now), align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    document.add_page_break()

    # ---------- 扉页：编制 / 审核 / 批准 ----------
    _para(document, report["title"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16)
    _para(document, project.name, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    for _ in range(4):
        document.add_paragraph()
    for label in ("编 制：", "审 核：", "批 准："):
        _para(document, f"{label}                    ", size=12)
        document.add_paragraph()
    _para(document, project.issuer, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    document.add_page_break()

    # ---------- 正文 ----------
    for block in report["blocks"]:
        kind = block["kind"]
        if kind == "h1":
            document.add_heading(block["text"], level=1)
        elif kind == "h2":
            document.add_heading(block["text"], level=2)
        elif kind == "p":
            _para(document, block["text"], bold=block.get("strong", False))
        elif kind == "caption":
            document.add_paragraph()
            _para(document, block["text"], align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
        elif kind == "table":
            rows = block["rows"]
            if block.get("info"):
                rows = rows + [
                    ["报告版本", f"V{version}"],
                    ["报告生成时间", f"{now:%Y-%m-%d %H:%M}"],
                    ["报告生成人", display_name(generated_by)],
                ]
            _add_table(document, block["headers"], rows, widths=block.get("widths"))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 生成与存档
# ---------------------------------------------------------------------------


def base_report_no(project, now=None):
    now = now or timezone.localtime()
    return project.report_no or f"ITEI-{now:%Y}-{str(project.id)[:6].upper()}"


@transaction.atomic
def generate_report(project, report_type, actor):
    validate_can_generate(project, report_type, actor)
    report = build_report(project, report_type)
    now = timezone.localtime()

    max_version = project.reports.filter(report_type=report_type).aggregate(v=Max("version"))["v"] or 0
    version = max_version + 1
    report_no = f"{base_report_no(project, now)}-{TYPE_SUFFIX[report_type]}"

    payload = _docx_bytes(report, report_no, version, actor)
    sha256 = hashlib.sha256(payload).hexdigest()
    filename = f"{report['title']}_{project.name}_V{version}.docx"
    levels = {part["sys"].code: part["achieved"] for part in report["parts"]}

    record = Report(
        project=project,
        report_type=report_type,
        version=version,
        report_no=report_no,
        filename=filename,
        sha256=sha256,
        trl_level=levels.get(TRL),
        mrl_level=levels.get(MRL),
        generated_by=actor,
    )
    record.file.save(filename, ContentFile(payload), save=True)

    if not project.report_no:
        project.report_no = base_report_no(project, now)
        project.save(update_fields=["report_no", "updated_at"])
    for part in report["parts"]:
        assessment = part["assessment"]
        if assessment.status != Assessment.STATUS_REPORTED:
            assessment.status = Assessment.STATUS_REPORTED
            assessment.save(update_fields=["status", "updated_at"])

    AuditLog.objects.create(
        actor=actor,
        action="report.generate",
        model_name="Report",
        object_id=str(record.id),
        object_repr=filename,
        project=project,
        after={"report_type": report_type, "version": version, "sha256": sha256, "levels": levels},
    )
    return record
